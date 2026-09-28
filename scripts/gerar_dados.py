"""
Converte a planilha de upload da MobiliDADOS (xlsx) nos dados que o site usa.

Uso:
    python scripts/gerar_dados.py [caminho/da/planilha.xlsx]

Sem argumento, usa a única planilha .xlsx da pasta dados/. Nomes e unidades dos
indicadores vêm de dados/catalogo-indicadores.csv (editável no Excel). Gera:
    src/data/*.json          -> lidos pelo Next.js no build
    public/dados/*.csv       -> arquivos de dados abertos para download no site

Antes de gerar, confere a estrutura da planilha. Se algo estiver errado, lista os
problemas e termina com erro, e a publicação automática não acontece (o site
continua na versão anterior). As mensagens citam só abas, colunas e contagens,
nunca o conteúdo das células, porque o log da publicação é público.
"""
import csv
import datetime
import json
import os
import re
import sys
import unicodedata

import openpyxl

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
PASTA_DADOS = os.path.join(RAIZ, 'dados')
SAIDA_JSON = os.path.join(RAIZ, 'src', 'data')
SAIDA_CSV = os.path.join(RAIZ, 'public', 'dados')
CATALOGO = os.path.join(PASTA_DADOS, 'catalogo-indicadores.csv')

# As 9 regiões metropolitanas monitoradas (sigla da planilha -> slug da URL)
RMS = {'RMB': 'rmb', 'RMBH': 'rmbh', 'RMC': 'rmc', 'RIDE-DF': 'ride-df', 'RMF': 'rmf',
       'RMR': 'rmr', 'RMRJ': 'rmrj', 'RMS': 'rms', 'RMSP': 'rmsp'}
NOME_CURTO = {'rmb': 'Belém', 'rmbh': 'Belo Horizonte', 'rmc': 'Curitiba',
              'ride-df': 'Distrito Federal e Entorno', 'rmf': 'Fortaleza', 'rmr': 'Recife',
              'rmrj': 'Rio de Janeiro', 'rms': 'Salvador', 'rmsp': 'São Paulo'}
MODOS_TMA = ['barca', 'brt', 'metro', 'monotrilho', 'trem', 'vlt']
# Colunas das abas de indicadores que identificam o local/ano (não são indicadores)
COLUNAS_ID = {'CD_UF', 'UF', 'CD_RM', 'RM', 'ANO', 'CD_MUN', 'COD_MUN6', 'CAPITAIS', 'CD_MUN_6'}

# Estrutura mínima que a planilha precisa ter
COLUNAS_OBRIGATORIAS = {
    'Info_gerais_capitais': {'CAPITAL', 'UF', 'TXT APRESENTACAO', 'Área total (km²)', 'Pop_2016', 'IDH-M'},
    'Info_gerais_RMs': {'SIGLA', 'RM', 'UF', 'TXT APRESENTACAO'},
    'Indicadores_capitais': {'CAPITAIS', 'ANO'},
    'Indicadores_RMs': {'RM', 'ANO'},
}
TOTAL_CAPITAIS = 27
ANO_MIN, ANO_MAX = 1990, 2100


class ErroPlanilha(Exception):
    """Problemas de estrutura que impedem a publicação."""


def num(v):
    """Converte célula em número; '-', vazio e textos viram None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if not isinstance(v, float):
        try:
            v = float(str(v).strip().replace(',', '.'))
        except ValueError:
            return None
    v = round(v, 6)  # remove ruído de ponto flutuante (0.33999999999999997 -> 0.34)
    return int(v) if v.is_integer() else v


def cel(v):
    """Valor para CSV em português: vírgula decimal, vazio para ausente."""
    if v is None:
        return ''
    return str(v).replace('.', ',') if isinstance(v, float) else v


def txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def slug(s):
    s = unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')


def ler(wb, aba):
    """Lê uma aba como lista de dicionários {cabeçalho: valor}, ignorando linhas vazias."""
    linhas = [r for r in wb[aba].iter_rows(values_only=True) if any(c not in (None, '') for c in r)]
    if not linhas:
        return []
    cab = [str(h).strip() if h is not None else None for h in linhas[0]]
    return [{cab[i]: r[i] for i in range(len(cab)) if cab[i]} for r in linhas[1:]]


def colunas(wb, aba):
    """Cabeçalhos de uma aba (mesmo que ela não tenha nenhuma linha de dados)."""
    for r in wb[aba].iter_rows(values_only=True):
        if any(c not in (None, '') for c in r):
            return {str(h).strip() for h in r if h is not None}
    return set()


def encontrar_planilha():
    """A única .xlsx da pasta dados/ (ignora arquivos temporários do Excel, que começam com ~$)."""
    xlsx = sorted(f for f in os.listdir(PASTA_DADOS) if f.lower().endswith('.xlsx') and not f.startswith('~$'))
    if len(xlsx) != 1:
        raise ErroPlanilha([f'A pasta dados/ deve ter exatamente 1 planilha .xlsx; encontrei {len(xlsx)}: '
                            + (', '.join(xlsx) or 'nenhuma') + '. Apague a antiga ao subir a nova.'])
    return os.path.join(PASTA_DADOS, xlsx[0])


def validar_estrutura(wb):
    """Abas e colunas obrigatórias presentes."""
    erros = []
    for aba, obrigatorias in COLUNAS_OBRIGATORIAS.items():
        if aba not in wb.sheetnames:
            erros.append(f'Aba "{aba}" não encontrada.')
            continue
        faltando = sorted(obrigatorias - colunas(wb, aba))
        if faltando:
            erros.append(f'Aba "{aba}": faltam as colunas ' + ', '.join(f'"{c}"' for c in faltando) + '.')
        elif not ler(wb, aba):
            erros.append(f'Aba "{aba}" não tem nenhuma linha de dados.')
    if erros:
        raise ErroPlanilha(erros)


def validar_conteudo(wb, caps, rms):
    """Locais e anos coerentes. Retorna avisos (que não impedem a publicação)."""
    erros, avisos = [], []
    if len(caps) != TOTAL_CAPITAIS:
        erros.append(f'Aba "Info_gerais_capitais": esperava {TOTAL_CAPITAIS} capitais, encontrei {len(caps)}.')
    slugs = [c['slug'] for c in caps]
    if len(set(slugs)) != len(slugs):
        erros.append('Aba "Info_gerais_capitais": há capitais repetidas.')
    if len(rms) != len(RMS):
        erros.append(f'Aba "Info_gerais_RMs": esperava as {len(RMS)} RMs monitoradas (coluna SIGLA), '
                     f'encontrei {len(rms)}.')

    for aba, coluna, lugares in (('Indicadores_capitais', 'CAPITAIS', caps), ('Indicadores_RMs', 'RM', rms)):
        nomes = {l['nome'] for l in lugares}
        linhas = ler(wb, aba)
        anos_invalidos = 0
        for r in linhas:
            ano = num(r.get('ANO'))
            if ano is None or not ANO_MIN <= ano <= ANO_MAX:
                anos_invalidos += 1
        if anos_invalidos:
            erros.append(f'Aba "{aba}": {anos_invalidos} linha(s) com ANO vazio ou fora de {ANO_MIN}-{ANO_MAX}.')
        sem_local = sum(1 for r in linhas if txt(r.get(coluna)) not in nomes)
        if sem_local:
            avisos.append(f'Aba "{aba}": {sem_local} linha(s) de locais fora da lista monitorada foram ignoradas.')
        # Textos em colunas de indicador (ex.: "NA", "n/d") são tratados como "sem dado"
        textos = {}
        for r in linhas:
            for k, v in r.items():
                if k not in COLUNAS_ID and v not in (None, '', '-') and num(v) is None:
                    textos[k] = textos.get(k, 0) + 1
        for k, n in sorted(textos.items()):
            avisos.append(f'Aba "{aba}", coluna "{k}": {n} célula(s) com texto em vez de número '
                          '(tratadas como sem dado).')
    if erros:
        raise ErroPlanilha(erros)
    return avisos


def data_da_planilha(wb):
    """Data em que a planilha foi salva pela última vez (ou hoje, se o arquivo não informar)."""
    modificada = wb.properties.modified or wb.properties.created
    return (modificada.date() if modificada else datetime.date.today()).isoformat()


def ler_rms(wb):
    rms = []
    for r in ler(wb, 'Info_gerais_RMs'):
        sig = txt(r.get('SIGLA')).upper()
        if sig not in RMS:
            continue
        s = RMS[sig]
        rms.append({
            'slug': s, 'sigla': sig, 'nome': txt(r['RM']), 'curto': NOME_CURTO[s], 'uf': txt(r['UF']),
            'lat': num(r['LAT']), 'lon': num(r['LONG']), 'texto': txt(r['TXT APRESENTACAO']),
            'area': num(r.get('Área total (km²)')), 'areaUrbana': num(r.get('Área urbana (km²)')),
            'pop2016': num(r.get('Pop_2016')), 'densidade': num(r.get('Densidade (hab./km²)')),
            'densidadeUrbana': num(r.get('Densidade urbana (hab./km²)')), 'idhm': num(r.get('IDHM')),
            'faixaIdhm': txt(r.get('Faixa do IDHM')), 'renda': num(r.get('Renda média')),
            'percDr1sm': num(r.get('Perc_DR_<1SM')), 'percNegros': num(r.get('Perc_Negros')),
            'percMulheres': num(r.get('Perc_Mulheres')),
        })
    rms.sort(key=lambda x: list(RMS.values()).index(x['slug']))
    return rms


def ler_capitais(wb):
    caps = []
    for r in ler(wb, 'Info_gerais_capitais'):
        nome = txt(r.get('CAPITAL'))
        if not nome:
            continue
        g = r.get
        caps.append({
            'slug': slug(nome), 'nome': nome, 'uf': txt(g('UF')), 'rm': txt(g('RM')),
            'lat': num(g('LAT')), 'lon': num(g('LONG')), 'texto': txt(g('TXT APRESENTACAO')),
            'area': num(g('Área total (km²)')), 'areaUrbana': num(g('Área urbana (km²)')),
            'pop2016': num(g('Pop_2016')), 'densidade': num(g('Densidade (hab./km²)')),
            'densidadeUrbana': num(g('Densidade urbana (hab./km²)')), 'idhm': num(g('IDH-M')),
            'faixaIdhm': txt(g('Faixa de IDH-M')), 'renda': num(g('Renda média')),
            'percDr1sm': num(g('Perc_DR_<1SM')), 'percNegros': num(g('Perc_Negros')),
            'percMulheres': num(g('Perc_Mulheres')), 'percBrancos': num(g('Perc_Brancos')),
            'percHomens': num(g('Perc_Homens')),
            'laiContrato': txt(g('lai_contrato')), 'laiContratoInicio': txt(g('lai_contrato_inicio')),
            'laiContratoPrazo': txt(g('lai_contrato_prazo')), 'laiContratoFonte': txt(g('lai_contrato_fonte')),
            'laiGps': txt(g('lai_gps')), 'laiGpsFrota': txt(g('lai_gps_frota')), 'laiGpsFonte': txt(g('lai_gps_fonte')),
            'laiGtfs': txt(g('lai_gtfs')), 'laiGtfsFonte': txt(g('lai_gtfs_fonte')),
            'tma': {m: {'estacoes': num(g('tma_estacao_' + m)), 'km': num(g('tma_km_' + m))} for m in MODOS_TMA},
            'planmobStatus': txt(g('planmob_status')), 'planmobAno': txt(g('planmob_ano_aprovacao')),
            'planmobFonte': txt(g('planmob_fonte')),
        })
    caps.sort(key=lambda c: slug(c['nome']))
    return caps


def ler_series(wb, aba, coluna_local, nome_para_slug):
    """Séries anuais no formato {indicador: {slug_do_local: {ano: valor}}}."""
    out = {}
    for r in ler(wb, aba):
        s = nome_para_slug.get(txt(r.get(coluna_local)))
        ano = num(r.get('ANO'))
        if not s or ano is None:
            continue
        for k, v in r.items():
            v = num(v)
            if k in COLUNAS_ID or v is None:
                continue
            out.setdefault(k, {}).setdefault(s, {})[str(int(ano))] = round(v, 4)
    return out


def salvar_json(nome, dados):
    with open(os.path.join(SAIDA_JSON, nome), 'w', encoding='utf-8', newline='\n') as f:
        json.dump(dados, f, ensure_ascii=False, indent=1)
        f.write('\n')


def salvar_csv(nome, linhas):
    # ';' e BOM para abrir direto no Excel em português
    with open(os.path.join(SAIDA_CSV, nome), 'w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f, delimiter=';').writerows(linhas)


def ler_catalogo():
    with open(CATALOGO, encoding='utf-8-sig', newline='') as f:
        return {r['codigo']: {'nome': r['nome'], 'unidade': r['unidade']} for r in csv.DictReader(f, delimiter=';')}


def csv_series(series, lugares, catalogo):
    nomes = {l['slug']: l['nome'] for l in lugares}
    linhas = [['local', 'ano', 'indicador', 'descricao', 'unidade', 'valor']]
    for k in sorted(series):
        info = catalogo.get(k, {'nome': k, 'unidade': ''})
        for s, anos in series[k].items():
            for ano, v in sorted(anos.items()):
                linhas.append([nomes[s], ano, k, info['nome'], info['unidade'], cel(v)])
    return linhas


def main(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    validar_estrutura(wb)
    rms = ler_rms(wb)
    caps = ler_capitais(wb)
    avisos = validar_conteudo(wb, caps, rms)
    ind_rms = ler_series(wb, 'Indicadores_RMs', 'RM', {r['nome']: r['slug'] for r in rms})
    ind_caps = ler_series(wb, 'Indicadores_capitais', 'CAPITAIS', {c['nome']: c['slug'] for c in caps})

    catalogo = ler_catalogo()
    sem_nome = sorted((set(ind_caps) | set(ind_rms)) - set(catalogo))
    if sem_nome:
        avisos.append('Indicadores sem nome em dados/catalogo-indicadores.csv (o site mostrará o código): '
                      + ', '.join(sem_nome))
    for a in avisos:
        print('AVISO:', a)

    os.makedirs(SAIDA_JSON, exist_ok=True)
    os.makedirs(SAIDA_CSV, exist_ok=True)
    salvar_json('meta.json', {'geradoEm': data_da_planilha(wb), 'planilha': os.path.basename(caminho)})
    salvar_json('catalogo-indicadores.json', catalogo)
    salvar_json('capitais.json', caps)
    salvar_json('indicadores-capitais.json', ind_caps)
    salvar_json('rms.json', rms)
    salvar_json('indicadores-rms.json', ind_rms)

    campos = [k for k in caps[0] if k != 'tma']
    tma = [f'tma_{t}_{m}' for m in MODOS_TMA for t in ('estacoes', 'km')]
    salvar_csv('capitais-informacoes-gerais.csv', [campos + tma] + [
        [cel(c[k]) for k in campos] + [cel(c['tma'][m][t]) for m in MODOS_TMA for t in ('estacoes', 'km')] for c in caps])
    salvar_csv('capitais-indicadores.csv', csv_series(ind_caps, caps, catalogo))
    campos_rm = list(rms[0])
    salvar_csv('rms-informacoes-gerais.csv', [campos_rm] + [[cel(r[k]) for k in campos_rm] for r in rms])
    salvar_csv('rms-indicadores.csv', csv_series(ind_rms, rms, catalogo))

    print(f'OK: {os.path.basename(caminho)} -> {len(caps)} capitais ({len(ind_caps)} indicadores), '
          f'{len(rms)} RMs ({len(ind_rms)} indicadores)')


if __name__ == '__main__':
    try:
        main(sys.argv[1] if len(sys.argv) > 1 else encontrar_planilha())
    except ErroPlanilha as e:
        print('ERRO: a planilha não passou na verificação. Nada foi publicado.')
        for problema in e.args[0]:
            print(' -', problema)
        sys.exit(1)
