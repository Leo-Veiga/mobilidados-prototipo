"""
Leva para a planilha do site os indicadores recalculados no repositório mobilidados-indicadores.

Uso:
    python scripts/importar_indicadores.py <pasta saida/ do mobilidados-indicadores> [planilha.xlsx]

Só atualiza colunas que a planilha já tem (o site ainda não aceita indicadores novos) e
acrescenta as linhas de anos novos de cada capital e RM. Duas regras:
    COMPARAVEIS  -> substitui os anos que o cálculo novo cobre; anos anteriores ficam como estão
    METODO_NOVO  -> apaga a coluna inteira e grava só a série nova (o método mudou)
Fórmulas da planilha viram valores (o conversor do site lê só valores, e o openpyxl não
recalcula fórmulas ao salvar).
"""
import csv
import glob
import os
import sys

import openpyxl

COMPARAVEIS = {
    'TX_MORT_TOT', 'TX_MORT_PED', 'TX_MORT_CICL', 'TX_MORT_AUTO', 'TX_MORT_MOTO', 'PERC_MORTE_NEGROS',
    'TX_INTERN_TOT', 'TX_INTERN_PED', 'TX_INTERN_CIC', 'TX_INTERN_AUTO', 'TX_INTERN_MOTO',
    'TX_MOTO', 'POP',
}
METODO_NOVO = {
    'PNT', 'PNT_ATE_1/2', 'PNT_1/2_1', 'PNT_1_3', 'PNT_ACIMA_3', 'PNT_MULHERES_NEGRAS', 'PNT_MULHERES_1SM',
    'PNB', 'PNB_ATE_1/2', 'PNB_1/2_1', 'PNB_1_3', 'PNB_ACIMA_3', 'PNB_MULHERES_NEGRAS', 'PNB_MULHERES_1SM',
    'PNB_MULHERES_ATE_2SM',  # não é mais calculado: fica vazio
    'TMA_INFRA_CICLO', 'TARIFAxSM',
}
ARQUIVOS = ['populacao', 'mortalidade', 'internacoes', 'motorizacao', 'pnb', 'pnt', 'tarifas', 'tma_ciclo']
# aba -> coluna com o código do local (código IBGE da capital ou CD_RM)
ABAS = {'Indicadores_capitais': 'CD_MUN', 'Indicadores_RMs': 'CD_RM'}


def ler_saida(pasta):
    """{indicador: {(local, ano): valor}} dos CSVs longos local;ano;indicador;valor."""
    dados = {}
    for nome in ARQUIVOS:
        with open(os.path.join(pasta, nome + '.csv'), encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                ind = r['indicador']
                if ind not in COMPARAVEIS | METODO_NOVO or r['valor'] in ('', 'NA'):
                    continue
                v = float(r['valor'])
                dados.setdefault(ind, {})[(int(r['local']), int(r['ano']))] = int(v) if v.is_integer() else v
    return dados


def congelar_formulas(wb, caminho):
    valores = openpyxl.load_workbook(caminho, data_only=True)
    for ws in wb:
        for linha in ws.iter_rows():
            for c in linha:
                if isinstance(c.value, str) and c.value.startswith('='):
                    c.value = valores[ws.title][c.coordinate].value


def atualizar_aba(ws, col_local, dados):
    cab = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    c_local, c_ano = cab[col_local], cab['ANO']
    linhas, ultima, modelo = {}, 1, {}
    for r in range(2, ws.max_row + 1):
        local, ano = ws.cell(r, c_local).value, ws.cell(r, c_ano).value
        if local is None or ano is None:
            continue
        linhas[(int(local), int(ano))] = r
        ultima = r
        modelo[int(local)] = r  # última linha do local: modelo para as colunas de identificação
    # Locais que o cálculo novo não cobre (ex.: RM de Porto Alegre) ficam intactos
    locais = set(modelo) & {l for serie in dados.values() for l, _ in serie}
    resumo = {}
    for ind in sorted((set(dados) | METODO_NOVO) & set(cab)):
        col = cab[ind]
        serie = {k: v for k, v in dados.get(ind, {}).items() if k[0] in locais}
        if not serie and ind not in METODO_NOVO:
            continue
        inicio = min((a for _, a in serie), default=0)
        apagadas = 0
        for (local, ano), r in linhas.items():
            if local in locais and (ind in METODO_NOVO or ano >= inicio):
                if ws.cell(r, col).value not in (None, '') and (local, ano) not in serie:
                    apagadas += 1
                ws.cell(r, col).value = None
        for (local, ano), v in sorted(serie.items()):
            if (local, ano) not in linhas:
                ultima += 1
                for c in range(1, ws.max_column + 1):
                    nome = ws.cell(1, c).value
                    if nome and str(nome).strip() in {'CD_UF', 'UF', 'CD_MUN', 'COD_MUN6', 'CAPITAIS',
                                                       'CD_RM', 'RM', 'CD_MUN_6'}:
                        ws.cell(ultima, c).value = ws.cell(modelo[local], c).value
                ws.cell(ultima, c_ano).value = ano
                linhas[(local, ano)] = ultima
            ws.cell(linhas[(local, ano)], col).value = v
        anos = sorted({a for _, a in serie}) or ['-']
        resumo[ind] = (anos[0], anos[-1], len(serie), apagadas)
    return resumo


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    pasta = sys.argv[1]
    caminho = sys.argv[2] if len(sys.argv) > 2 else glob.glob(
        os.path.join(os.path.dirname(__file__), '..', 'dados', '*.xlsx'))[0]
    dados = ler_saida(pasta)
    wb = openpyxl.load_workbook(caminho)
    congelar_formulas(wb, caminho)
    for aba, col_local in ABAS.items():
        print(aba)
        for ind, (a0, a1, n, apagadas) in sorted(atualizar_aba(wb[aba], col_local, dados).items()):
            extra = f'; {apagadas} valor(es) antigo(s) sem correspondente apagado(s)' if apagadas else ''
            print(f'  {ind:24} {a0}-{a1}  {n} valores{extra}')
        sem_coluna = sorted(set(dados) - {str(c.value).strip() for c in wb[aba][1] if c.value})
        if sem_coluna:
            print('  sem coluna na aba (ignorados):', ', '.join(sem_coluna))
    wb.save(caminho)


if __name__ == '__main__':
    main()
